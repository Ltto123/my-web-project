"""词表的确定性读取：保留源字段和重复词义，不把学习统计交给 AI。"""
import csv
import io
import json
import re
import zipfile


TEXT_FIELDS = ("pos", "def_en", "def_zh", "example_en", "example_zh")
MAX_SOURCE_ENTRIES = 5000
MAX_SOURCE_CHARS = 1_000_000
MAX_SHEET_ROWS = 20000
MAX_SHEET_COLUMNS = 64
_ALIASES = {
    "word": {"word", "words", "vocabulary", "单词", "词汇", "英文单词", "英文", "词语", "单词短语"},
    "pos": {"pos", "partofspeech", "词性"},
    "def_zh": {"defzh", "中文", "中文释义", "中文意思"},
    "meaning": {"meaning", "definition", "释义", "含义"},
    "translation": {"translation", "翻译"},
    "def_en": {"defen", "englishdefinition", "englishmeaning", "英文释义", "英文解释", "英语释义"},
    "example_en": {"exampleen", "example", "sentence", "英文例句", "英语例句", "例句"},
    "example_zh": {"examplezh", "exampletranslation", "sentencetranslation", "中文例句", "例句翻译", "例句中文翻译"},
}


def _header(value):
    return re.sub(r"[\s_\-/（）()]+", "", str(value or "")).casefold()


def _columns(row):
    mapped = {}
    for index, value in enumerate(row):
        name = _header(value)
        for field, aliases in _ALIASES.items():
            if name in aliases:
                if field in mapped:
                    # 并排的两块词表不能只取左侧，交给未知布局识别。
                    return None
                mapped[field] = index
    return mapped if "word" in mapped else None


def _text(value):
    return value.strip() if isinstance(value, str) else ""


def merge_entries(entries: list[dict]) -> list[dict]:
    """相同单词按第一次出现排序；不同源义项逐项保留。"""
    merged = {}
    variants = {}
    source_chars = 0
    if len(entries) > MAX_SHEET_ROWS:
        raise ValueError("原始词条过多，请将文件拆小后上传")
    for row in entries:
        if not isinstance(row, dict):
            raise ValueError("词条必须是包含 word 的对象")
        word = _text(row.get("word"))
        if not word or len(word) > 200 or not re.search(r"[A-Za-z]", word):
            raise ValueError("存在空白或无效的英文词条，请检查单词列")
        key = re.sub(r"\s+", " ", word).casefold()
        if key not in merged:
            if len(merged) >= MAX_SOURCE_ENTRIES:
                raise ValueError(f"单份文件最多支持 {MAX_SOURCE_ENTRIES} 个不同词条")
            merged[key] = {"word": word, "is_phrase": int(bool(re.search(r"\s", word)))}
            variants[key] = {field: [] for field in TEXT_FIELDS}
        for field in TEXT_FIELDS:
            value = _text(row.get(field))
            source_chars += len(value)
            if source_chars > MAX_SOURCE_CHARS:
                raise ValueError("词表文字内容过长，请将文件拆小后上传")
            if value and value not in variants[key][field]:
                variants[key][field].append(value)
    for key, row in merged.items():
        for field in TEXT_FIELDS:
            row[field] = "\n".join(variants[key][field]) or None
    return list(merged.values())


def _rows_to_entries(rows):
    mapping = None
    entries = []
    nonempty = 0
    for row in rows:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        nonempty += 1
        if mapping is None:
            mapping = _columns(row)
            if mapping is None and nonempty >= 20:
                return None
            continue
        # 分页复制的表头不是词条。
        if _columns(row) == mapping:
            continue
        entry = {field: row[index] if index < len(row) else None for field, index in mapping.items()}
        for neutral in ("meaning", "translation"):
            value = _text(entry.pop(neutral, None))
            if value:
                field = "def_zh" if re.search(r"[\u4e00-\u9fff]", value) else "def_en"
                existing = _text(entry.get(field))
                entry[field] = existing + "\n" + value if existing and existing != value else value
        if not _text(entry.get("word")):
            # 有其他词典内容的无单词行不能悄悄丢弃。
            if any(_text(entry.get(field)) for field in TEXT_FIELDS):
                raise ValueError("词表中有释义但单词为空，请检查原文件")
            continue
        entries.append(entry)
    return entries if mapping is not None or not nonempty else None


def extract_xlsx_entries(contents: bytes) -> list[dict] | None:
    """识别有表头的 Excel；未知布局返回 None，由调用方走文本识别。"""
    from openpyxl import load_workbook

    # 压缩文件体积不能代表展开后的内存成本。
    with zipfile.ZipFile(io.BytesIO(contents)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("Excel 展开后的内容过大，请精简工作表后上传")
    book = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    all_entries = []
    try:
        # 先检查全部工作表，不能因封面布局未知而跳过后续表的资源限制。
        for sheet in book.worksheets:
            if sheet.max_row and sheet.max_row > MAX_SHEET_ROWS:
                raise ValueError("Excel 工作表行数过多，请移除空白格式行或拆分文件")
            if sheet.max_column and sheet.max_column > MAX_SHEET_COLUMNS:
                raise ValueError("Excel 列数过多，请仅保留词表相关列")
        for sheet in book.worksheets:
            entries = _rows_to_entries(sheet.iter_rows(values_only=True))
            if entries is None:
                return None
            all_entries.extend(entries)
    finally:
        book.close()
    return merge_entries(all_entries) if all_entries else None


def extract_obvious_line(line: str) -> dict | None:
    """保守识别带编号的单词行/双语行，供确定性抽取及 AI 漏词校验使用。"""
    line = line.strip()
    # 多个并排词条不能当作「一个单词 + 后面所有内容」；交给布局识别。
    for delimiter in ("\t", ",", ";"):
        if delimiter in line:
            fields = next(csv.reader([line], delimiter=delimiter))
            if sum(bool(field.strip()) for field in fields) > 2:
                return None
    unnumbered = re.sub(r"^(?:\d+[.)、]|[（(]\d+[）)])\s*", "", line)
    if unnumbered != line and re.fullmatch(r"[A-Za-z][A-Za-z'’.-]*", unnumbered):
        return {"word": unnumbered}
    match = re.fullmatch(
        r"([A-Za-z][A-Za-z'’ .()/&-]{0,150}?)\s*"
        r"(?:[\t:：,，]|[—–]+|\s+[-]+\s+|\s+(?=[\u4e00-\u9fff]))\s*(.+)", unnumbered)
    if not match or not re.search(r"[\u4e00-\u9fff]", match[2]):
        return None
    return {"word": match[1].strip(), "def_zh": match[2].strip()}


def extract_text_entries(content: str) -> list[dict] | None:
    """结构明确的 JSON、带表头分隔文本及逐行单词无需 AI 猜测词表。"""
    content = content.strip()
    if not content:
        return []
    if len(content) > MAX_SOURCE_CHARS:
        raise ValueError("文件文字内容过长，请拆分后上传")
    if content.startswith(("[", "{")):
        try:
            decoded = json.loads(content)
            if isinstance(decoded, dict):
                decoded = decoded.get("words")
            if isinstance(decoded, list) and all(isinstance(row, dict) and "word" in row for row in decoded):
                return merge_entries(decoded)
        except json.JSONDecodeError:
            pass
    for delimiter in ("\t", ",", ";"):
        if delimiter not in content.splitlines()[0]:
            continue
        rows = list(csv.reader(io.StringIO(content), delimiter=delimiter))
        if rows and _columns(rows[0]):
            entries = _rows_to_entries(rows)
            if entries is not None:
                return merge_entries(entries)
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    if all(re.fullmatch(r"[A-Za-z][A-Za-z'’.-]*", line) for line in lines):
        return merge_entries([{"word": line} for line in lines])
    # 最常见的无表头双语单词表；必须每一行都符合才走确定性路径。
    entries = []
    for line in lines:
        entry = extract_obvious_line(line)
        if entry is None:
            return None
        entries.append(entry)
    return merge_entries(entries)
